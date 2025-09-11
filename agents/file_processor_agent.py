"""
文件处理Agent
支持多种文件格式的解析和数据提取
"""
import os
import pandas as pd
import PyPDF2
import docx
from PIL import Image
import pytesseract
import openpyxl
from typing import Dict, Any, Optional, List
import json
import logging
from pathlib import Path

class FileProcessorAgent:
    def __init__(self):
        self.supported_formats = {
            '.pdf': self._process_pdf,
            '.docx': self._process_docx,
            '.doc': self._process_doc,
            '.xlsx': self._process_excel,
            '.xls': self._process_excel,
            '.csv': self._process_csv,
            '.txt': self._process_txt,
            '.png': self._process_image,
            '.jpg': self._process_image,
            '.jpeg': self._process_image,
            '.json': self._process_json
        }
        self.logger = logging.getLogger(__name__)
        
    def process_file(self, file_path: str) -> Dict[str, Any]:
        """
        处理上传的文件，返回标准化的数据结构
        """
        try:
            file_path = Path(file_path)
            if not file_path.exists():
                raise FileNotFoundError(f"文件不存在: {file_path}")
            
            file_ext = file_path.suffix.lower()
            if file_ext not in self.supported_formats:
                raise ValueError(f"不支持的文件格式: {file_ext}")
            
            # 调用对应的处理函数
            processor = self.supported_formats[file_ext]
            result = processor(str(file_path))
            
            # 添加元数据
            result['metadata'] = {
                'filename': file_path.name,
                'file_size': file_path.stat().st_size,
                'file_type': file_ext,
                'processed_at': pd.Timestamp.now().isoformat()
            }
            
            return result
            
        except Exception as e:
            self.logger.error(f"文件处理失败: {e}")
            return {
                'success': False,
                'error': str(e),
                'data': None,
                'data_type': 'error'
            }
    
    def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """处理PDF文件"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text_content = []
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text_content.append(page.extract_text())
                
                full_text = '\n'.join(text_content)
                
                # 尝试从文本中提取表格数据
                tables = self._extract_tables_from_text(full_text)
                
                return {
                    'success': True,
                    'data': {
                        'text': full_text,
                        'tables': tables,
                        'pages': len(pdf_reader.pages)
                    },
                    'data_type': 'document_with_tables' if tables else 'document'
                }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """处理Word文档"""
        try:
            doc = docx.Document(file_path)
            text_content = []
            tables_data = []
            
            # 提取文本
            for paragraph in doc.paragraphs:
                text_content.append(paragraph.text)
            
            # 提取表格
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                tables_data.append(table_data)
            
            return {
                'success': True,
                'data': {
                    'text': '\n'.join(text_content),
                    'tables': tables_data
                },
                'data_type': 'document_with_tables' if tables_data else 'document'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_doc(self, file_path: str) -> Dict[str, Any]:
        """处理旧版Word文档（需要额外工具）"""
        # 这里可以使用python-docx2txt或其他工具
        try:
            # 简化处理，建议用户转换为docx格式
            return {
                'success': False,
                'error': '请将.doc文件转换为.docx格式后重新上传',
                'data_type': 'error'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_excel(self, file_path: str) -> Dict[str, Any]:
        """处理Excel文件"""
        try:
            # 读取所有工作表
            excel_file = pd.ExcelFile(file_path)
            sheets_data = {}
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                sheets_data[sheet_name] = {
                    'data': df.to_dict('records'),
                    'columns': df.columns.tolist(),
                    'shape': df.shape,
                    'dtypes': df.dtypes.to_dict()
                }
            
            return {
                'success': True,
                'data': sheets_data,
                'data_type': 'structured_data'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_csv(self, file_path: str) -> Dict[str, Any]:
        """处理CSV文件"""
        try:
            # 尝试不同的编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']
            df = None
            
            for encoding in encodings:
                try:
                    df = pd.read_csv(file_path, encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
            
            if df is None:
                raise ValueError("无法解码CSV文件")
            
            return {
                'success': True,
                'data': {
                    'data': df.to_dict('records'),
                    'columns': df.columns.tolist(),
                    'shape': df.shape,
                    'dtypes': df.dtypes.to_dict()
                },
                'data_type': 'structured_data'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_txt(self, file_path: str) -> Dict[str, Any]:
        """处理文本文件"""
        try:
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']
            content = None
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue
            
            if content is None:
                raise ValueError("无法解码文本文件")
            
            # 尝试从文本中提取结构化数据
            tables = self._extract_tables_from_text(content)
            
            return {
                'success': True,
                'data': {
                    'text': content,
                    'tables': tables
                },
                'data_type': 'document_with_tables' if tables else 'document'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_image(self, file_path: str) -> Dict[str, Any]:
        """处理图片文件，使用OCR提取文本"""
        try:
            # 打开图片
            image = Image.open(file_path)
            
            # 使用OCR提取文字
            text = pytesseract.image_to_string(image, lang='chi_sim+eng')
            
            # 尝试从提取的文本中识别表格数据
            tables = self._extract_tables_from_text(text)
            
            return {
                'success': True,
                'data': {
                    'text': text,
                    'tables': tables,
                    'image_size': image.size
                },
                'data_type': 'ocr_document_with_tables' if tables else 'ocr_document'
            }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _process_json(self, file_path: str) -> Dict[str, Any]:
        """处理JSON文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            # 如果是列表且包含字典，转换为DataFrame格式
            if isinstance(data, list) and data and isinstance(data[0], dict):
                df = pd.DataFrame(data)
                return {
                    'success': True,
                    'data': {
                        'data': df.to_dict('records'),
                        'columns': df.columns.tolist(),
                        'shape': df.shape,
                        'dtypes': df.dtypes.to_dict()
                    },
                    'data_type': 'structured_data'
                }
            else:
                return {
                    'success': True,
                    'data': data,
                    'data_type': 'json_data'
                }
        except Exception as e:
            return {'success': False, 'error': str(e), 'data_type': 'error'}
    
    def _extract_tables_from_text(self, text: str) -> List[List[List[str]]]:
        """从文本中提取表格数据"""
        tables = []
        lines = text.split('\n')
        
        # 简单的表格识别算法
        current_table = []
        for line in lines:
            line = line.strip()
            if not line:
                if current_table:
                    tables.append(current_table)
                    current_table = []
                continue
            
            # 检查是否包含表格分隔符
            if '\t' in line or '|' in line or ',' in line:
                if '\t' in line:
                    row = [cell.strip() for cell in line.split('\t')]
                elif '|' in line:
                    row = [cell.strip() for cell in line.split('|') if cell.strip()]
                else:
                    row = [cell.strip() for cell in line.split(',')]
                
                if len(row) > 1:  # 至少有两列才认为是表格
                    current_table.append(row)
        
        if current_table:
            tables.append(current_table)
        
        return tables

if __name__ == "__main__":
    # 测试代码
    agent = FileProcessorAgent()
    print("文件处理Agent初始化完成")
    print(f"支持的文件格式: {list(agent.supported_formats.keys())}")